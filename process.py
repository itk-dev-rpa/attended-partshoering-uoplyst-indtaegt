"""This module contains the main process of the robot."""

from datetime import datetime
import time
from dataclasses import dataclass
from tkinter import filedialog, messagebox

from selenium.webdriver.chrome.webdriver import WebDriver as Chrome
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.wait import WebDriverWait
from selenium.common.exceptions import InvalidSessionIdException
from openpyxl import load_workbook


@dataclass
class BankInfo:
    """A dataclass representing information about a bank."""
    name: str
    reg_number: str
    transactions: list[tuple[datetime, str, int]]


def process():
    """The main process of the robot."""
    excel_path = filedialog.askopenfilename()

    if not excel_path:
        messagebox.showinfo("Annulleret", "Ingen fil valgt. Robotten stopper.")
        return

    bank_info = read_excel(excel_path)

    if not show_bank_info_popup(bank_info):
        messagebox.showinfo("Annulleret", "Bankoplysninger ikke korrekt. Robotten stopper.")
        return

    browser = login()

    for i, bank in enumerate(bank_info):
        insert_bank(bank, i+1, browser)

        if len(bank_info) > i+1:
            add_bank(i+1, browser)

    messagebox.showinfo("Robot færdig", "Bankdata indsat.\nHusk selv at udfylde CPR-nummer øverst og indsende formularen.\nLuk browseren når du er færdig.")
    wait_for_browser_close(browser)
    messagebox.showinfo("Browser lukket", "Browseren er lukket. Robotten er færdig.")


def login():
    """Open Chrome and navigate to the form."""
    browser = Chrome()
    browser.implicitly_wait(2)
    browser.maximize_window()
    browser.get("https://selvbetjening.aarhuskommune.dk/da/content/partshoering-vedroerende-potentiel-uoplyst-indtaegt-12")

    return browser


def read_excel(excel_path: str) -> list[BankInfo]:
    """Read the given Excel file.
    Each sheet in the file is treated as a separate bank.

    Args:
        excel_path: The path to the Excel file.

    Returns:
        A list for BankInfo objects.
    """
    wb = load_workbook(excel_path, read_only=True)

    banks = []

    for ws in wb.worksheets:
        values = []

        bank_name = ws["A2"].value
        bank_reg = ws["C2"].value

        for row in ws:
            if isinstance(row[0].value, datetime):
                values.append([row[0].value.strftime("%Y-%m-%d"), row[1].value, row[2].value])

        banks.append(BankInfo(bank_name, bank_reg, values[0:20]))

    return banks


def insert_bank(bank_info: BankInfo, index: int, browser: Chrome):
    """Insert info about a bank into the form.

    Args:
        bank_info: The bank info to insert.
        index: The index of the bank in the form.
        browser: The browser to perform the action.
    """
    # Fill out bank data
    browser.find_element(By.NAME, f"bankens_navn{index}").send_keys(bank_info.name)
    browser.find_element(By.NAME, f"bankens_registreringsnummer{index}").send_keys(bank_info.reg_number)

    # Add lines to custom composite
    add_lines(browser, index, len(bank_info.transactions)-1)

    # Fill out rows using javascript
    js = """
        const values = arguments[0];
        for (let i = 0; i < values.length; i++) {{
            document.getElementsByName(`transaktioner{index}[items][${i}][dato{index}]`)[0].value = values[i][0];
            document.getElementsByName(`transaktioner{index}[items][${i}][tekst{index}]`)[0].value = values[i][1];
            document.getElementsByName(`transaktioner{index}[items][${i}][beloeb{index}]`)[0].value = values[i][2];
        }}
        """.replace("{index}", str(index))

    browser.execute_script(js, bank_info.transactions)


def add_lines(browser: Chrome, index: int, count: int):
    """Add the given number of lines to the custom composite in the form.
    OS2Forms only allows adding 100 lines at a time, so the
    function inserts 100 lines multiple times until the desired count is reached.

    Args:
        browser: The browser to perform the action.
        index: The index of the bank in the form.
        count: The number of lines to add.
    """
    wait = WebDriverWait(browser, 10)

    while count > 0:
        add_button = wait.until(EC.element_to_be_clickable((By.NAME, f"transaktioner{index}_table_add")))
        count_input = browser.find_element(By.NAME, f"transaktioner{index}[add][more_items]")
        count_input.clear()
        count_input.send_keys(min(count, 100))
        add_button.click()
        count -= 100
        time.sleep(1)

    wait.until(EC.element_to_be_clickable((By.NAME, f"transaktioner{index}_table_add")))


def add_bank(index: int, browser: Chrome):
    """Check the 'Tilføj Bank' of the given bank index."""
    browser.find_element(By.NAME, f"tilfoej_bank{index}").click()


def show_bank_info_popup(bank_info: list[BankInfo]) -> bool:
    """Show a popup summarizing the bank info and
    ask for user confirmation.

    Args:
        bank_info: Bank info to summarize.

    Returns:
        True if the user accepts the info else false.
    """
    message = f"Der blev fundet {len(bank_info)} banker i filen:\n\n"

    for bank in bank_info:
        message += f"Banknavn: {bank.name}\nBank reg: {bank.reg_number}\nAntal transaktioner: {len(bank.transactions)}\n\n"

    message += "Er dette korrekt?"

    return messagebox.askyesno("Bankinfo", message)


def wait_for_browser_close(browser: Chrome, timeout: int = 1200):
    """Wait for the browser window to close.
    If the timeout is reached the browser is closed automatically.
    """
    try:
        while timeout > 0:
            browser.current_window_handle  # pylint: disable=pointless-statement
            time.sleep(1)
            timeout -= 1
    except InvalidSessionIdException:
        ...

    if timeout == 0:
        browser.quit()


if __name__ == '__main__':
    process()
